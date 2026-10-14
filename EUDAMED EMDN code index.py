#!/usr/bin/env python3
"""EUDAMED Market Intelligence Tool

Features:
- Search a local EMDN Excel dictionary with one or more keywords.
- Use the verified EUDAMED cndCode filter.
- Compress parent/child EMDN matches to avoid downloading overlapping branches.
- Optional manufacturer filter, including manufacturer-only searches.
- Paginate strictly according to totalPages.
- Retry transient HTTP failures.
- Export Devices, Manufacturer Ranking, AR Ranking, EMDN Summary,
  Matched EMDN Codes, and Search Criteria sheets.
- Automatically name the output file from the active filters.
"""

import math
import re
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ============================================================================
# USER SETTINGS: normally you only need to edit this section
# ============================================================================

LOOKUP_EMDN_KEYWORDS = [
    "blood sugar",
    "blood glucose",
]

# Leave blank when you do not want to restrict by manufacturer.
SEARCH_MANUFACTURER = ""

#replace the path below with the path to your EMDN dictionary file. The default is set to a local path on the user's machine.
EMDN_DICTIONARY = Path(
    "/Users/<username>/EMDN v2026_EN..xlsx"
)

# True is recommended. Example: if R06 and R0601 both match, only R06 is
# downloaded because R06 already covers the R0601 branch in EUDAMED.
COMPRESS_PARENT_CHILD_CODES = True

# EUDAMED currently returns up to 50 rows per page in this workflow.
PAGE_SIZE = 50
REQUEST_DELAY_SECONDS = 0.5
OUTPUT_DIRECTORY = Path.home() / "Desktop"

# Set to None for no safety cap. A positive integer limits pages per code.
MAX_PAGES_PER_CODE = None

# ============================================================================
# CONSTANTS
# ============================================================================

URL = "https://ec.europa.eu/tools/eudamed/api/devices/udiDiData"
STATUS_ON_MARKET = "refdata.device-model-status.on-the-market"
EMDN_TERM_COLUMN = "TERM EMDN (v.2026)"
EMDN_CODE_COLUMN = "CODE"

DEVICE_COLUMNS = [
    "EMDN Code",
    "EMDN Description",
    "Primary UDI-DI / EUDAMED ID",
    "Basic UDI-DI",
    "Trade Name",
    "Risk Class",
    "Manufacturer",
    "Authorised Representative",
]

# ============================================================================
# TEXT AND RECORD HELPERS
# ============================================================================


def clean_text(value):
    """Return a trimmed string and safely ignore None/dicts/numbers."""
    return value.strip() if isinstance(value, str) else ""


def nested_name(node):
    """Extract an organisation name from common nested API structures."""
    if isinstance(node, dict):
        return (
            clean_text(node.get("organisationName"))
            or clean_text(node.get("name"))
            or clean_text(node.get("legalName"))
        )
    return clean_text(node)


def parse_risk_class(risk):
    """Read either a risk-class object or a risk-class code string."""
    if isinstance(risk, dict):
        label = clean_text(risk.get("name"))
        code = clean_text(risk.get("code"))
    else:
        label = ""
        code = clean_text(risk)

    if label:
        return label
    if not code:
        return ""

    value = code.replace("refdata.risk-class.", "").replace("-", " ").title()
    replacements = {
        "Ivd": "IVD",
        "Iii": "III",
        "Iib": "IIb",
        "Iia": "IIa",
        "Ii": "II",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value


def extract_basic_udi(item):
    """Extract Basic UDI-DI from known response shapes."""
    basic = item.get("basicUdi") or item.get("basicUdiData")
    if isinstance(basic, dict):
        value = clean_text(basic.get("basicUdiDi"))
    else:
        value = clean_text(basic)
    return value or clean_text(item.get("basicUdiDi"))


def extract_record(item, emdn_code, emdn_description):
    """Convert one EUDAMED item to a flat Excel row."""
    manufacturer = (
        clean_text(item.get("manufacturerName"))
        or clean_text(item.get("producerName"))
        or clean_text(item.get("systemProducerName"))
        or nested_name(item.get("manufacturer"))
        or nested_name(item.get("producer"))
        or nested_name(item.get("systemProducer"))
    )

    representative = (
        clean_text(item.get("authorisedRepresentativeName"))
        or nested_name(item.get("authorisedRepresentative"))
        or nested_name(item.get("authRep"))
    )

    return {
        "EMDN Code": emdn_code,
        "EMDN Description": emdn_description,
        "Primary UDI-DI / EUDAMED ID": (
            clean_text(item.get("primaryDi"))
            or clean_text(item.get("eudamedId"))
        ),
        "Basic UDI-DI": extract_basic_udi(item),
        "Trade Name": clean_text(item.get("tradeName")),
        "Risk Class": parse_risk_class(item.get("riskClass")),
        "Manufacturer": manufacturer,
        "Authorised Representative": representative,
    }

# ============================================================================
# HTTP SESSION
# ============================================================================


def make_session():
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            "Accept": "application/json, text/plain, */*",
        }
    )

    retry = Retry(
        total=5,
        connect=5,
        read=5,
        status=5,
        backoff_factor=1.0,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=frozenset(["GET"]),
        respect_retry_after_header=True,
        raise_on_status=False,
    )
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session

# ============================================================================
# EMDN DICTIONARY
# ============================================================================


def validate_settings():
    if PAGE_SIZE < 1:
        raise ValueError("PAGE_SIZE must be at least 1.")
    if MAX_PAGES_PER_CODE is not None and MAX_PAGES_PER_CODE < 1:
        raise ValueError("MAX_PAGES_PER_CODE must be None or at least 1.")
    if not LOOKUP_EMDN_KEYWORDS and not SEARCH_MANUFACTURER.strip():
        raise ValueError(
            "Enter at least one EMDN keyword or a manufacturer name."
        )
    if LOOKUP_EMDN_KEYWORDS and not EMDN_DICTIONARY.exists():
        raise FileNotFoundError(f"EMDN dictionary not found: {EMDN_DICTIONARY}")


def read_emdn_dictionary(filepath):
    """Read the dictionary and tolerate either row 1 or row 2 as the header."""
    last_error = None
    for header_row in (1, 0):
        try:
            df = pd.read_excel(filepath, header=header_row, engine="openpyxl")
            if EMDN_CODE_COLUMN in df.columns and EMDN_TERM_COLUMN in df.columns:
                return df
        except Exception as exc:
            last_error = exc

    raise ValueError(
        f"Could not find columns '{EMDN_CODE_COLUMN}' and "
        f"'{EMDN_TERM_COLUMN}' in {filepath}. Last error: {last_error}"
    )


def search_emdn_dictionary(keywords, filepath):
    """Match literal keyword phrases using OR logic, just like keyword lookup."""
    if not keywords:
        return pd.DataFrame(columns=[EMDN_CODE_COLUMN, EMDN_TERM_COLUMN])

    cleaned_keywords = [str(k).strip() for k in keywords if str(k).strip()]
    if not cleaned_keywords:
        return pd.DataFrame(columns=[EMDN_CODE_COLUMN, EMDN_TERM_COLUMN])

    dictionary = read_emdn_dictionary(filepath)
    literal_pattern = "|".join(re.escape(k) for k in cleaned_keywords)
    term_series = dictionary[EMDN_TERM_COLUMN].fillna("").astype(str)

    matches = dictionary.loc[
        term_series.str.contains(literal_pattern, case=False, regex=True),
        [EMDN_CODE_COLUMN, EMDN_TERM_COLUMN],
    ].copy()

    matches[EMDN_CODE_COLUMN] = matches[EMDN_CODE_COLUMN].astype(str).str.strip()
    matches[EMDN_TERM_COLUMN] = matches[EMDN_TERM_COLUMN].astype(str).str.strip()
    matches = matches[
        matches[EMDN_CODE_COLUMN].ne("")
        & matches[EMDN_CODE_COLUMN].ne("nan")
    ]
    matches = matches.drop_duplicates(subset=[EMDN_CODE_COLUMN]).reset_index(drop=True)

    print(f"\nFound {len(matches)} matching EMDN codes:")
    for _, row in matches.iterrows():
        print(f"{row[EMDN_CODE_COLUMN]} -> {row[EMDN_TERM_COLUMN]}")
    return matches


def compress_parent_child_codes(matches):
    """Keep the shortest matched ancestor in each EMDN code branch."""
    if matches.empty:
        return matches.copy(), matches.copy()

    ordered = matches.copy()
    ordered["_length"] = ordered[EMDN_CODE_COLUMN].str.len()
    ordered = ordered.sort_values(["_length", EMDN_CODE_COLUMN])

    kept_codes = []
    removed_rows = []
    kept_rows = []

    for _, row in ordered.iterrows():
        code = row[EMDN_CODE_COLUMN]
        parent = next((p for p in kept_codes if code != p and code.startswith(p)), None)
        row_dict = {
            EMDN_CODE_COLUMN: code,
            EMDN_TERM_COLUMN: row[EMDN_TERM_COLUMN],
        }
        if parent:
            row_dict["Covered By Parent"] = parent
            removed_rows.append(row_dict)
        else:
            kept_codes.append(code)
            kept_rows.append(row_dict)

    kept = pd.DataFrame(kept_rows, columns=[EMDN_CODE_COLUMN, EMDN_TERM_COLUMN])
    removed = pd.DataFrame(
        removed_rows,
        columns=[EMDN_CODE_COLUMN, EMDN_TERM_COLUMN, "Covered By Parent"],
    )
    return kept.reset_index(drop=True), removed.reset_index(drop=True)

# ============================================================================
# EUDAMED REQUESTS
# ============================================================================


def build_params(page, code="", manufacturer_filter=""):
    params = {
        "deviceStatusCode": STATUS_ON_MARKET,
        "page": page,
        "size": PAGE_SIZE,
        "sort": "primaryDi,desc",
        "languageIsoCode": "en",
    }
    if code:
        params["cndCode"] = code
    if manufacturer_filter.strip():
        params["manufacturerName"] = manufacturer_filter.strip()
    return params


def request_page(session, page, code="", manufacturer_filter=""):
    params = build_params(page, code, manufacturer_filter)
    response = session.get(URL, params=params, timeout=(15, 60))
    response.raise_for_status()
    data = response.json()
    if not isinstance(data, dict) or not isinstance(data.get("content"), list):
        raise ValueError(
            "Unexpected EUDAMED response: expected a JSON object with a content list."
        )
    return data


def inspect_target(session, code="", description="", manufacturer_filter=""):
    """Fetch page 1 once and use it both for workload estimation and download."""
    data = request_page(session, 0, code, manufacturer_filter)
    total_records = data.get("totalElements")
    api_total_pages = data.get("totalPages")

    if not isinstance(total_records, int):
        total_records = len(data.get("content", []))

    if isinstance(api_total_pages, int):
        total_pages = api_total_pages
    else:
        total_pages = math.ceil(total_records / PAGE_SIZE) if total_records else 0

    return {
        "EMDN Code": code,
        "EMDN Description": description,
        "Records": total_records,
        "Pages": total_pages,
        "first_page": data,
    }


def fetch_target(session, target, manufacturer_filter=""):
    code = target["EMDN Code"]
    description = target["EMDN Description"]
    total_records = target["Records"]
    total_pages = target["Pages"]

    display_name = code if code else f"Manufacturer: {manufacturer_filter}"
    print(f"\n===== {display_name} =====")
    print(f"Total records: {total_records}")
    print(f"Total pages: {total_pages}")

    pages_to_fetch = total_pages
    if MAX_PAGES_PER_CODE is not None:
        pages_to_fetch = min(pages_to_fetch, MAX_PAGES_PER_CODE)

    records = []
    for page in range(pages_to_fetch):
        if page == 0:
            data = target["first_page"]
        else:
            data = request_page(session, page, code, manufacturer_filter)

        items = data.get("content", [])
        print(f"Page {page + 1}/{pages_to_fetch} ({len(items)} rows)")

        for item in items:
            if isinstance(item, dict):
                records.append(extract_record(item, code, description))

        if not items:
            print("No rows returned; stopping this target early.")
            break
        if page + 1 < pages_to_fetch:
            time.sleep(REQUEST_DELAY_SECONDS)

    return records

# ============================================================================
# DATA CLEANING AND SUMMARIES
# ============================================================================


def normalize_blank_columns(df):
    for column in DEVICE_COLUMNS:
        if column not in df.columns:
            df[column] = ""
        df[column] = df[column].fillna("").astype(str).str.strip()
    return df[DEVICE_COLUMNS].copy()


def deduplicate_devices(df):
    """Deduplicate safely and retain all matching EMDN codes/descriptions."""
    if df.empty:
        return normalize_blank_columns(df), 0

    df = normalize_blank_columns(df)

    def identity(row):
        primary = row["Primary UDI-DI / EUDAMED ID"]
        if primary:
            return "UDI::" + primary.casefold()
        fallback = "||".join(
            [
                row["Basic UDI-DI"],
                row["Trade Name"],
                row["Manufacturer"],
            ]
        ).casefold()
        return "FALLBACK::" + fallback

    working = df.copy()
    working["_identity"] = working.apply(identity, axis=1)

    def join_unique(values):
        unique = []
        for value in values:
            for part in str(value).split("; "):
                part = part.strip()
                if part and part not in unique:
                    unique.append(part)
        return "; ".join(unique)

    aggregations = {
        "EMDN Code": join_unique,
        "EMDN Description": join_unique,
        "Primary UDI-DI / EUDAMED ID": "first",
        "Basic UDI-DI": "first",
        "Trade Name": "first",
        "Risk Class": "first",
        "Manufacturer": "first",
        "Authorised Representative": "first",
    }

    deduped = (
        working.groupby("_identity", sort=False, as_index=False)
        .agg(aggregations)
        .drop(columns=["_identity"], errors="ignore")
    )
    deduped = deduped[DEVICE_COLUMNS]
    removed = len(df) - len(deduped)
    return deduped, removed


def make_ranking(df, column):
    valid = df[df[column].fillna("").astype(str).str.strip().ne("")].copy()
    if valid.empty:
        return pd.DataFrame(columns=[column, "Unique Device Count"])
    return (
        valid.groupby(column, dropna=False)
        .size()
        .reset_index(name="Unique Device Count")
        .sort_values(["Unique Device Count", column], ascending=[False, True])
        .reset_index(drop=True)
    )

# ============================================================================
# EXCEL OUTPUT
# ============================================================================


def safe_filename_part(value, max_length=60):
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value).strip())
    value = re.sub(r"_+", "_", value).strip("_.")
    return (value or "search")[:max_length]


def build_output_path():
    parts = []
    if LOOKUP_EMDN_KEYWORDS:
        parts.append(safe_filename_part("_".join(LOOKUP_EMDN_KEYWORDS)))
    if SEARCH_MANUFACTURER.strip():
        parts.append("MFG_" + safe_filename_part(SEARCH_MANUFACTURER))
    parts.append(datetime.now().strftime("%Y%m%d_%H%M"))

    output_dir = OUTPUT_DIRECTORY.expanduser()
    if not output_dir.exists():
        output_dir = Path.cwd()
    return output_dir / ("EUDAMED_" + "_".join(parts) + ".xlsx")


def style_workbook(writer):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        for column_cells in worksheet.columns:
            letter = get_column_letter(column_cells[0].column)
            maximum = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[letter].width = min(max(maximum + 2, 12), 60)


def export_excel(
    devices_df,
    emdn_summary_df,
    all_matches_df,
    removed_children_df,
    raw_rows,
    duplicates_removed,
    output_file,
):
    manufacturer_ranking = make_ranking(devices_df, "Manufacturer")
    ar_ranking = make_ranking(devices_df, "Authorised Representative")

    criteria_rows = [
        ("Search keywords", ", ".join(LOOKUP_EMDN_KEYWORDS)),
        ("Manufacturer filter", SEARCH_MANUFACTURER),
        ("Dictionary path", str(EMDN_DICTIONARY)),
        ("Parent-child compression", COMPRESS_PARENT_CHILD_CODES),
        ("Page size", PAGE_SIZE),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Raw downloaded rows", raw_rows),
        ("Duplicates removed", duplicates_removed),
        ("Final unique devices", len(devices_df)),
    ]
    criteria_df = pd.DataFrame(criteria_rows, columns=["Setting", "Value"])

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        devices_df.to_excel(writer, sheet_name="Devices", index=False)
        manufacturer_ranking.to_excel(
            writer, sheet_name="Manufacturer Ranking", index=False
        )
        ar_ranking.to_excel(writer, sheet_name="AR Ranking", index=False)
        emdn_summary_df.to_excel(writer, sheet_name="EMDN Summary", index=False)
        all_matches_df.to_excel(writer, sheet_name="Matched EMDN Codes", index=False)
        if not removed_children_df.empty:
            removed_children_df.to_excel(
                writer, sheet_name="Skipped Child Codes", index=False
            )
        criteria_df.to_excel(writer, sheet_name="Search Criteria", index=False)
        style_workbook(writer)

# ============================================================================
# MAIN
# ============================================================================


def main():
    validate_settings()

    all_matches = search_emdn_dictionary(
        LOOKUP_EMDN_KEYWORDS, EMDN_DICTIONARY
    )

    if LOOKUP_EMDN_KEYWORDS and all_matches.empty:
        print("No EMDN codes matched the supplied keywords. No file was created.")
        return

    if COMPRESS_PARENT_CHILD_CODES and not all_matches.empty:
        selected_codes, removed_children = compress_parent_child_codes(all_matches)
    else:
        selected_codes = all_matches.copy()
        removed_children = pd.DataFrame(
            columns=[EMDN_CODE_COLUMN, EMDN_TERM_COLUMN, "Covered By Parent"]
        )

    if not selected_codes.empty:
        print("\nCodes selected for downloading:")
        for _, row in selected_codes.iterrows():
            print(f"{row[EMDN_CODE_COLUMN]} -> {row[EMDN_TERM_COLUMN]}")

    if not removed_children.empty:
        print("\nChild codes skipped because a matched parent covers the branch:")
        for _, row in removed_children.iterrows():
            print(
                f"{row[EMDN_CODE_COLUMN]} -> covered by "
                f"{row['Covered By Parent']}"
            )

    # Manufacturer-only mode uses one target with no cndCode.
    targets_to_inspect = []
    if selected_codes.empty and SEARCH_MANUFACTURER.strip():
        targets_to_inspect.append(("", "Manufacturer-only search"))
    else:
        for _, row in selected_codes.iterrows():
            targets_to_inspect.append(
                (str(row[EMDN_CODE_COLUMN]), str(row[EMDN_TERM_COLUMN]))
            )

    all_records = []
    inspected_targets = []

    with make_session() as session:
        print("\nEstimating workload...")
        for code, description in targets_to_inspect:
            target = inspect_target(
                session,
                code=code,
                description=description,
                manufacturer_filter=SEARCH_MANUFACTURER,
            )
            inspected_targets.append(target)
            label = code or f"Manufacturer: {SEARCH_MANUFACTURER}"
            print(
                f"{label}: {target['Records']} records, "
                f"{target['Pages']} pages"
            )

        estimated_pages = sum(t["Pages"] for t in inspected_targets)
        if MAX_PAGES_PER_CODE is not None:
            estimated_pages = sum(
                min(t["Pages"], MAX_PAGES_PER_CODE) for t in inspected_targets
            )
        print(f"Estimated pages to download: {estimated_pages}")

        for target in inspected_targets:
            records = fetch_target(
                session,
                target,
                manufacturer_filter=SEARCH_MANUFACTURER,
            )
            all_records.extend(records)

    raw_df = pd.DataFrame(all_records, columns=DEVICE_COLUMNS)
    raw_rows = len(raw_df)
    devices_df, duplicates_removed = deduplicate_devices(raw_df)

    summary_rows = []
    for target in inspected_targets:
        summary_rows.append(
            {
                "EMDN Code": target["EMDN Code"],
                "EMDN Description": target["EMDN Description"],
                "API Records": target["Records"],
                "API Pages": target["Pages"],
            }
        )
    emdn_summary_df = pd.DataFrame(
        summary_rows,
        columns=["EMDN Code", "EMDN Description", "API Records", "API Pages"],
    )

    output_file = build_output_path()
    export_excel(
        devices_df=devices_df,
        emdn_summary_df=emdn_summary_df,
        all_matches_df=all_matches,
        removed_children_df=removed_children,
        raw_rows=raw_rows,
        duplicates_removed=duplicates_removed,
        output_file=output_file,
    )

    print("\n==============================")
    print("Completed")
    print("==============================")
    print(f"Raw downloaded rows: {raw_rows}")
    print(f"Duplicates removed: {duplicates_removed}")
    print(f"Final unique devices: {len(devices_df)}")
    print(f"Excel saved to: {output_file}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nStopped by user. No incomplete Excel file was written.")
    except PermissionError as exc:
        print(
            "\nCould not write the Excel file. Close any open file with the same "
            f"name and try again. Details: {exc}"
        )
    except requests.RequestException as exc:
        print(f"\nEUDAMED request failed: {exc}")
    except Exception as exc:
        print(f"\nError: {type(exc).__name__}: {exc}")
